"""Canonical text extraction for supported text-based source formats."""

from __future__ import annotations

from pathlib import Path
from typing import Any


class TextExtractor:
    """Extract normalized text payloads from supported file types.

    Returns a canonical dict with:
    - ``full_text``: concatenated plain text
    - ``segments``: structure-aware segment list
    """

    SUPPORTED_FORMATS = {"pdf", "docx", "epub", "txt", "pptx"}

    def extract(self, path: str | Path, file_type: str) -> dict[str, Any]:
        """Dispatch extraction by normalized file type."""
        normalized_type = file_type.lower()
        if normalized_type not in self.SUPPORTED_FORMATS:
            raise ValueError(f"Unsupported file type: {file_type}")

        handler = getattr(self, f"_extract_{normalized_type}")
        return handler(Path(path))

    def _extract_pdf(self, path: Path) -> dict[str, Any]:
        """Extract page-aware text from a PDF file."""
        import fitz

        segments: list[dict[str, Any]] = []
        full_text_parts: list[str] = []
        with fitz.open(path) as document:
            for page_index, page in enumerate(document, start=1):
                text = page.get_text("text").strip()
                if not text:
                    continue
                segments.append({"page": page_index, "text": text})
                full_text_parts.append(text)
        return {"segments": segments, "full_text": "\n\n".join(full_text_parts)}

    def _extract_docx(self, path: Path) -> dict[str, Any]:
        """Extract heading-aware paragraph groups from a DOCX file."""
        from docx import Document

        document = Document(path)
        segments: list[dict[str, Any]] = []
        full_text_parts: list[str] = []
        current_heading = ""
        buffer: list[str] = []

        def flush_buffer() -> None:
            nonlocal buffer
            text = "\n".join(part for part in buffer if part.strip()).strip()
            if not text:
                buffer = []
                return
            segment = {"heading": current_heading, "text": text}
            segments.append(segment)
            full_text_parts.append(text)
            buffer = []

        for paragraph in document.paragraphs:
            text = paragraph.text.strip()
            if not text:
                continue
            style_name = getattr(paragraph.style, "name", "") or ""
            if style_name.lower().startswith("heading"):
                flush_buffer()
                current_heading = text
                continue
            buffer.append(text)

        flush_buffer()
        return {"segments": segments, "full_text": "\n\n".join(full_text_parts)}

    def _extract_epub(self, path: Path) -> dict[str, Any]:
        """Extract chapter-aware plain text from an EPUB file."""
        from bs4 import BeautifulSoup
        from ebooklib import ITEM_DOCUMENT, epub

        book = epub.read_epub(str(path))
        segments: list[dict[str, Any]] = []
        full_text_parts: list[str] = []

        for chapter_index, item in enumerate(book.get_items_of_type(ITEM_DOCUMENT), start=1):
            soup = BeautifulSoup(item.get_body_content(), "html.parser")
            text = soup.get_text("\n", strip=True).strip()
            if not text:
                continue
            segment = {"chapter": chapter_index, "title": item.get_name(), "text": text}
            segments.append(segment)
            full_text_parts.append(text)

        return {"segments": segments, "full_text": "\n\n".join(full_text_parts)}

    def _extract_txt(self, path: Path) -> dict[str, Any]:
        """Extract plain text and paragraph-based segments from a TXT file."""
        full_text = path.read_text(encoding="utf-8").strip()
        paragraphs = [part.strip() for part in full_text.split("\n\n") if part.strip()]
        segments = [{"paragraph": index, "text": text} for index, text in enumerate(paragraphs, start=1)]
        return {"segments": segments, "full_text": full_text}

    def _extract_pptx(self, path: Path) -> dict[str, Any]:
        """Extract slide-aware text from a PPTX file."""
        from pptx import Presentation

        presentation = Presentation(path)
        segments: list[dict[str, Any]] = []
        full_text_parts: list[str] = []
        for slide_index, slide in enumerate(presentation.slides, start=1):
            texts = [
                shape.text.strip()
                for shape in slide.shapes
                if getattr(shape, "has_text_frame", False) and shape.text.strip()
            ]
            if not texts:
                continue
            slide_text = "\n".join(texts)
            segments.append({"slide": slide_index, "text": slide_text})
            full_text_parts.append(slide_text)
        return {"segments": segments, "full_text": "\n\n".join(full_text_parts)}
